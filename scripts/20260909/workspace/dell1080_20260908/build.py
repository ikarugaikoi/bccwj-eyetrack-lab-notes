'DELL1080 migration through existing, user-authorized guest UI executor.\nAll native project changes use Pro Lab UI; native evidence is read-only.\n'
import argparse, hashlib, importlib.util, json, re, shutil, subprocess, sys, time, traceback, zipfile
from pathlib import Path
ROOT=Path(__file__).resolve().parent
AUTO=ROOT.parent/'local_automation'
sys.path.insert(0,str(AUTO))
from controller import command, BRIDGE
from ui import UI,find
from production import open_editor,set_panel,verify_container_ui,descendants
from prolab import submit_dialog as original_submit_dialog
from runner import read_design,normalize,write
from batch import exact_dialog
from acceptance import policy,permitted,RECORD
PACKAGE=ROOT.parent/'review_dell1080_20260908/BCCWJ_TOBII_FULL_20260908_DELL1080_56CM_120HZ'
GUEST='__MASKED_LOCAL_PATH_0291__'
TEST='__MASKED_LOCAL_PATH_0292__'
CAT={lang:json.loads((PACKAGE/f'04_PROJECT_BLUEPRINTS/PROJECTS_{lang}.json').read_text(encoding='utf-8')) for lang in ('ZH','JA')}
SPEC=json.loads((PACKAGE/'05_QA/EXPECTED_PROJECTS_QA_ONLY.json').read_text(encoding='utf-8'))
sp=importlib.util.spec_from_file_location('dell_checker',PACKAGE/'05_QA/check_native_project.py')
checker=importlib.util.module_from_spec(sp);sp.loader.exec_module(checker)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def queue():return json.loads((ROOT/'queue.json').read_text(encoding='utf-8'))
def bp(job):return next(p for p in CAT[job['language']]['projects'] if p['bilingual_project_name']==job['planned'])
def log(job,step,**extra):
    row=dict(time=time.time(),project=job['actual'],sequence=job['sequence'],step=step,**extra)
    with (ROOT/'logs/steps.jsonl').open('a',encoding='utf-8') as f:f.write(json.dumps(row,ensure_ascii=False)+'\n')
    write(ROOT/'logs/current.json',row)
    write(AUTO/'logs/job_status.json',dict(job='production_build',project=job['actual'],status=extra.get('status','running'),projects_verified=queue()['static_verified'],project_total=84,current=step,cloud_calls=0,error=extra.get('error','')))
    print(json.dumps(row,ensure_ascii=True),flush=True)
def stage(job,key,**extras):
    q=queue();j=next(j for j in q['jobs'] if j['actual']==job['actual']);j.update(stage=key,**extras)
    if key=='verified':j['status']='static_verified'
    q['static_verified']=sum(j['status']=='static_verified' for j in q['jobs']);q['remaining']=84-q['static_verified'];write(ROOT/'queue.json',q)
    job.update(stage=key,**extras)
def main_ui():
    ws=[w for w in command('windows') if w['title']=='Tobii Pro Lab']
    if len(ws)!=1:raise RuntimeError('Expected one guest Pro Lab window')
    command('activate',hwnd=ws[0]['hwnd']);return UI(ws[0]['hwnd'])

def submit_dialog(d,main_hwnd,button='DialogOkButton'):
    try:return original_submit_dialog(d,main_hwnd,button)
    except RuntimeError as exc:
        #  The guest can report a COM error in its post-action observation when
        #  a native dialog has just been destroyed. Never repeat the submission.
        if '-2146233083' not in str(exc):raise
        if any(w['hwnd']==d.hwnd for w in command('windows')):raise
        for attempt in range(4):
            try:
                recovered=UI(main_hwnd)
                assert recovered.s['foreground']==main_hwnd,'Unexpected new modal after dialog close'
                return recovered
            except RuntimeError as read_error:
                if '-2146233083' not in str(read_error) or attempt==3:raise
                time.sleep(.5)
def browse(u,path):
    if find(u.s,automation_id='ApplicationMenuHeader').get('selected')!=1:u.act('click',{'automation_id':'ApplicationMenuHeader'})
    if not any(e['automation_id']=='browseButton' and e['visible'] for e in u.s['elements']):u.act('click',{'automation_id':'tbHeader','name':'Open existing Project','type':'Text'})
    u.act('invoke',{'automation_id':'browseButton'})
    d=exact_dialog('Browse for project folder');d.act('set_value',{'automation_id':'1148','type':'Edit'},value=path)
    return submit_dialog(d,u.hwnd,'1')
def open_job(u,job):
    u=browse(u,job['guest_path']+'\\tobii.project')
    assert normalize(u.s.get('project',''))==normalize(job['actual'])
    u.act('click',{'automation_id':'designerNavigationButton'})
    if any(e['automation_id']=='BackToStructure' and e['visible'] for e in u.s['elements']):u.act('invoke',{'automation_id':'BackToStructure'})
    return u
def rename_job(u,job):
    expected=job['planned'] if job['language']=='ZH' else job['actual'][:-3]+'_ZH'
    log(job,'Open independent design copy')
    u=browse(u,job['source_path']+'\\tobii.project')
    assert normalize(u.s.get('project',''))==normalize(expected),(u.s.get('project'),expected)
    u=browse(u,TEST)
    u.act('click',{'automation_id':'ApplicationMenuHeader'})
    u.act('click',{'type':'ListItem','name':job['source_path']+'\\tobii.project'},right=True)
    u.act('click',{'type':'MenuItem','name':'Rename Project'})
    u.act('set_value',{'automation_id':'txtname'},value=job['actual'])
    u.act('key',key='ENTER')
    stage(job,'renamed');log(job,'Renamed to short project name')
    return open_job(u,job)
def resolution(u,job):
    d,_,_=read_design(job['guest_path'])
    values=[v for v in d['Settings'].values() if isinstance(v,str) and re.fullmatch('\\d+,\\d+',v)]
    if values==['1920,1080']:return u
    assert values==['1920,1200'],values
    combos=[e for e in u.s['elements'] if e['visible'] and e['type']=='ComboBox' and 'DisplayResolution' in str(e.get('selected_text'))]
    assert len(combos)==1
    u.act('click',{'id':combos[0]['id']})
    candidates=[e for e in u.s['elements'] if e['visible'] and e['enabled'] and e['type']=='Text' and e['depth']==3 and re.match('\\s*1920 x 1080\\b',e.get('name',''))]
    if len(candidates)!=1:
        write(ROOT/'logs/resolution_picker.json',u.s);raise RuntimeError('Inspect resolution choices: '+str([(e['type'],e['name']) for e in candidates]))
    u.act('click',{'id':candidates[0]['id']})
    d,_,_=read_design(job['guest_path']);assert '1920,1080' in d['Settings'].values()
    log(job,'Presentation resolution 1920x1080 verified');return u
def verify_table(job):
    import openpyxl
    b=bp(job);d,n,snapshot=read_design(job['guest_path']);assert len(d['Tables'])==1
    tid,t=next(iter(d['Tables'].items()));book=PACKAGE/b['workbook_path'];assert sha(book)==b['workbook_sha256']
    wb=openpyxl.load_workbook(book,read_only=True,data_only=True);rows=list(wb['Tobii_Design'].values);wb.close()
    assert n[tid]=='Tobii_Design' and [n[c] for c in t['Columns']]==list(rows[0])
    norm=lambda v:None if v is None else str(v)
    expected=[norm(v) for row in rows[1:] for v in row]
    assert len(t['Cells'])==len(expected)
    differences=[dict(excel_row=i//len(rows[0])+2,column=rows[0][i%len(rows[0])],observed=norm(a),expected=e) for i,(a,e) in enumerate(zip(t['Cells'],expected)) if norm(a)!=e]
    report=dict(strict_passed=not differences,accepted_passed=permitted(differences),acceptance=policy(),workbook=str(book),workbook_sha256=sha(book),differences=differences)
    write(ROOT/'qa'/(job['actual']+'_table_deviation.json'),report)
    assert report['accepted_passed'],'Unexpected table differences beyond approved DOT rounding'
    return d,n,snapshot
def update_table(u,job):
    before,_,_=read_design(job['guest_path'])
    u.act('click',{'automation_id':'ImportTableButton','type':'MenuItem'})
    u.act('click',{'automation_id':'UpdateTableMenuItem','type':'MenuItem','depth':2})
    d=exact_dialog('Update selected Design table')
    d.act('set_value',{'automation_id':'1148','type':'Edit'},value=GUEST+'\\Materials\\'+bp(job)['workbook_path'].replace('/','\\'))
    u=submit_dialog(d,u.hwnd,'1');u=submit_dialog(exact_dialog('UpdateTableWarning'),u.hwnd)
    after,_,_=verify_table(job)
    assert after['Bindings']==before['Bindings'] and after['Tables'].keys()==before['Tables'].keys()
    log(job,'Updated sealed table; all cells and unchanged bindings verified');return u
def numeric(u,aid,value):
    if float(find(u.s,automation_id=aid,type='Edit')['value'])!=value:
        u.act('set_value',{'automation_id':aid,'type':'Edit'},value=str(value));u.act('key',key='TAB')
    assert float(find(u.s,automation_id=aid,type='Edit')['value'])==value
def fixed_pages(u,job):
    evidence=[]
    for e in bp(job)['elements']:
        if e['kind']!='image_stimulus':continue
        open_editor(u,e['stimulus_name']);set_panel(u,'StimulusPropertiesTool',False)
        containers=[x for x in u.s['elements'] if x['automation_id']=='containerName' and x['visible']]
        assert len(containers)==1
        u.act('click',{'id':containers[0]['id']})
        numeric(u,'ContainerHeight',1);numeric(u,'ContainerY',0)
        assert float(find(u.s,automation_id='ContainerWidth',type='Edit')['value'])==1 and float(find(u.s,automation_id='ContainerX',type='Edit')['value'])==0
        result=verify_container_ui(u);result['stimulus']=e['stimulus_name'];evidence.append(result)
        u.act('invoke',{'automation_id':'BackToStructure'})
        log(job,'Fixed page geometry: '+e['stimulus_name'])
    write(ROOT/'qa'/(job['actual']+'_fixed_ui.json'),evidence)
    stage(job,job.get('stage','renamed'),fixed_pages_verified=True)
    return u
def source_for_job(job):
    if job['language']=='ZH':return BRIDGE/'native_backups'/(job['planned']+'.zip')
    return ROOT/'backups'/(job['actual'][:-3]+'_ZH.zip')
def read_backup_design(archive):
    with zipfile.ZipFile(archive) as z:
        members=[n for n in z.namelist() if '/Current/Design/' in n and n.endswith('.json')]
        assert len(members)==1
        return json.loads(z.read(members[0]))['Data']
def inherited_check(job,current):
    before=read_backup_design(source_for_job(job));expected=json.loads(json.dumps(before))
    #  This key tracks the selected top-level editor object. Confirmed by selecting
    #  PRACTICE_10_END: the key changed from absent to its exact native object ID.
    #  Evidence: fafe2d4fd3594d34ae6f37d8b687b623.zip vs 7e59968b23d342e092d2124a952a7604.zip.
    selection_key='i6872b519-5118-433b-97d8-0a1436dc58e5'+before['RootGroupId']
    selection_changes=[]
    if before['Settings'].get(selection_key)!=current['Settings'].get(selection_key):
        allowed={None,*before['StructureItems']}
        assert before['Settings'].get(selection_key) in allowed and current['Settings'].get(selection_key) in allowed
        selection_changes.append(dict(key=selection_key,before=before['Settings'].get(selection_key),after=current['Settings'].get(selection_key)))
        expected['Settings'].pop(selection_key,None)
        if selection_key in current['Settings']:expected['Settings'][selection_key]=current['Settings'][selection_key]
    #  Change only the named geometry values; compare all remaining native design fields.
    for k,v in expected['Settings'].items():
        if v=='1920,1200':expected['Settings'][k]='1920,1080'
    for tid,t in expected['Tables'].items():
        #  Entire final table is verified independently against the sealed workbook.
        assert t['Columns']==current['Tables'][tid]['Columns']
        t['Cells']=current['Tables'][tid]['Cells']
    for sid,item in expected['StructureItems'].items():
        if item['kind']!=5:continue
        c=item['data']
        for field,value in [('Height',1),('Y',0)]:
            prop=c.get(field)
            if prop and prop.get('Id') not in expected['Bindings']:
                prop['FixedValue']=value
    differences=[]
    def walk(a,b,p=''):
        if isinstance(a,dict) and isinstance(b,dict):
            for k in a.keys()|b.keys():
                if k not in a or k not in b:differences.append(p+'/'+k)
                else:walk(a[k],b[k],p+'/'+k)
        elif isinstance(a,list) and isinstance(b,list) and len(a)==len(b):
            for i,(x,y) in enumerate(zip(a,b)):walk(x,y,p+'/'+str(i))
        elif a!=b:differences.append(p)
    walk(expected,current)
    result={'source_backup_sha256':sha(source_for_job(job)),'editor_selection_changes':selection_changes,'differences_after_allowed_geometry':differences}
    write(ROOT/'qa'/(job['actual']+'_inheritance.json'),result)
    if differences:raise RuntimeError('Unexpected inherited design changes: '+str(differences[:20]))
def audit_snapshot(job):
    d,n,archive=verify_table(job)
    dest=ROOT/'snapshots'/job['actual'];dest.mkdir(exist_ok=True)
    with zipfile.ZipFile(archive) as z:
        for name in z.namelist():assert (dest/name).resolve().is_relative_to(dest.resolve())
        z.extractall(dest)
        names={Path(x).parts[0] for x in z.namelist()};assert len(names)==1
    result=checker.inspect(dest/next(iter(names)),job['planned'],SPEC)
    write(ROOT/'qa'/(job['actual']+'_native.json'),result)
    approved_spec=json.loads(json.dumps(SPEC))
    approved_spec['geometry']['DOT'].update(H=0.0778,Y=0.0667)
    accepted=checker.inspect(dest/next(iter(names)),job['planned'],approved_spec)
    accepted.update(acceptance=policy(),strict_failure_count=result['failure_count'],status='STATIC_CHECKS_PASSED_WITH_APPROVED_ROUNDING_NOT_HARDWARE_VERIFIED' if not accepted['failure_count'] else 'STATIC_CHECKS_FAILED')
    write(ROOT/'qa'/(job['actual']+'_native_accepted.json'),accepted)
    if accepted['failure_count']:raise RuntimeError('Native checks beyond approved deviation failed: '+str(accepted['failures'][:3]))
    if job['language']=='ZH':inherited_check(job,d)
    return accepted
def export(u,job,review=False):
    export_name=job['actual']+('_REVIEW' if review else '')
    target=BRIDGE/'native_backups'/(export_name+'.zip');assert not target.exists()
    u.act('click',{'automation_id':'DashboardRadioButton'})
    u.act('click',{'automation_id':'Export','type':'MenuItem'});u.act('click',{'automation_id':'Project','type':'MenuItem','depth':2})
    d=exact_dialog('ExportProjectDialogWindowCaption')
    if find(d.s,automation_id='ExportTypeComboboxId').get('selected_text')!='Full project':d.act('select_combo',{'automation_id':'ExportTypeComboboxId'},value='Full project')
    folder='__MASKED_LOCAL_PATH_0293__'
    if find(d.s,automation_id='projectExportFileLocationFieldId')['name']!=folder:
        d.act('invoke',{'automation_id':'projectExportSelectFileButtonId'});picker=exact_dialog('Project export')
        picker.act('set_value',{'automation_id':'1152','type':'Edit'},value=folder);d=submit_dialog(picker,d.hwnd,'1')
    d.act('set_value',{'automation_id':'projectExportFileNameFieldId'},value=export_name);d.act('invoke',{'automation_id':'projectExportExportButtonId'})
    for _ in range(40):
        done=any(e['automation_id']=='ExportDoneShowInFolderButton' and e['visible'] for e in d.s['elements'])
        close_ready=any(e['automation_id']=='DialogCancelButton' and e['type']=='Button' and e['visible'] and e['enabled'] for e in d.s['elements'])
        if done and close_ready:break
        time.sleep(.5);d.observe()
    else:raise RuntimeError('Export completion not observed')
    u=submit_dialog(d,u.hwnd,'DialogCancelButton')
    with zipfile.ZipFile(target) as z:assert z.testzip() is None and not any('/Recordings/' in n for n in z.namelist() if not n.endswith('/'))
    shutil.copy2(target,ROOT/'backups'/target.name)
    audit_args=[sys.executable,str(ROOT/'audit_full_backup.py'),str(target),'--package',str(PACKAGE),'--expected',job['planned'],'--actual',job['actual'],'--report',str(ROOT/'qa'/(export_name+'_full_backup.json'))]
    if not review:audit_args+=['--accept-record',str(RECORD)]
    subprocess.run(audit_args,check=not review)
    return u,target
def japanese_media(u,job):
    d,n,_=read_design(job['guest_path'])
    expected=[p.stem for p in (PACKAGE/'01_MEDIA/instructions_JA').glob('*.png')]
    if len(d['MediaLibrary'])==131:
        assert all(name in n.values() for name in expected);return u
    assert len(d['MediaLibrary'])==112
    set_panel(u,'MediaLibrary',True)
    u.act('invoke',{'automation_id':'SwitchableViewListBoxImportButton','type':'Button'})
    picker=exact_dialog('Import media into Media Library')
    picker.act('set_value',{'automation_id':'1148','type':'Edit'},value=GUEST+'\\Materials\\01_MEDIA\\instructions_JA')
    picker.act('invoke',{'automation_id':'1','type':'Button'})
    picker.act('click',{'type':'ListItem','name':'FORMAL_00_intro_JA'});picker.act('key',key='CTRL_A')
    selected=[e['name'] for e in picker.s['elements'] if e['type']=='ListItem' and e.get('selected')==1]
    assert sorted(selected)==sorted(expected)
    u=submit_dialog(picker,u.hwnd,'1');set_panel(u,'MediaLibrary',False)
    d,n,_=read_design(job['guest_path']);assert len(d['MediaLibrary'])==131 and all(name in n.values() for name in expected)
    log(job,'All 19 Japanese instruction media imported');return u
def japanese_pages(u,job):
    evidence=[]
    def current_source(change):
        design,names,_=read_design(job['guest_path'])
        ids=[k for k,v in names.items() if v==change['stimulus_name'] and k in design['StructureItems']]
        assert len(ids)==1
        cids=design['HierarchyInfo'][ids[0]]['ChildrenIdentities'];assert len(cids)==1
        return names[design['StructureItems'][cids[0]]['data']['MediaId']['FixedValue']]
    for change in job['blueprint']['instruction_replacements']:
        source_before=current_source(change)
        assert source_before in (change['from_source'],change['to_source']),'Unexpected existing instruction source'
        open_editor(u,change['stimulus_name']);set_panel(u,'StimulusPropertiesTool',False)
        containers=[e for e in u.s['elements'] if e['automation_id']=='containerName' and e['visible']];assert len(containers)==1
        u.act('click',{'id':containers[0]['id']})
        if source_before!=change['to_source']:
            p=find(u.s,automation_id='ContainerPropertiesTool')['id']
            u.act('invoke',{'automation_id':'AddMedia','parent':p})
            picker=exact_dialog('Add image');picker.act('set_value',{'automation_id':'1148','type':'Edit'},value=GUEST+'\\Materials\\'+change['to_media_path'].replace('/','\\'))
            u=submit_dialog(picker,u.hwnd,'1')
        assert current_source(change)==change['to_source'],'Japanese source was not saved'
        if find(u.s,automation_id='ContainerScaleOriginal').get('selected')!=1:u.act('click',{'automation_id':'ContainerScaleOriginal'})
        result=verify_container_ui(u);result.update(stimulus=change['stimulus_name'],source=change['to_source']);evidence.append(result)
        for aid,value in [('ContainerWidth',1),('ContainerHeight',1),('ContainerX',0),('ContainerY',0)]:assert float(find(u.s,automation_id=aid,type='Edit')['value'])==value
        u.act('invoke',{'automation_id':'BackToStructure'});log(job,'Japanese source: '+change['stimulus_name'])
    write(ROOT/'qa'/(job['actual']+'_fixed_ui.json'),evidence);return u
def build_ja(u,job):
    assert all(j['status']=='static_verified' for j in queue()['jobs'] if j['language']=='ZH')
    if job.get('stage') is None:u=rename_job(u,job)
    else:u=open_job(u,job)
    if job.get('stage') not in ('configured','reopened'):
        u=japanese_media(u,job);u=japanese_pages(u,job);stage(job,'configured')
    u=browse(u,TEST);u=open_job(u,job);stage(job,'reopened')
    audit=audit_snapshot(job);log(job,'Japanese native checks passed',checks=audit['check_count'])
    u,target=export(u,job);stage(job,'verified',backup=str(ROOT/'backups'/target.name),backup_sha256=sha(target),checks=audit['check_count'])
    log(job,'Project verified and full native backup complete',status='project_verified',completed=queue()['static_verified']);return u
def build_zh(u,job):
    if job.get('stage') is None:u=rename_job(u,job)
    else:u=open_job(u,job)
    u=resolution(u,job)
    if job.get('stage') not in ('configured','reopened'):
        u=update_table(u,job)
        if not job.get('fixed_pages_verified'):u=fixed_pages(u,job)
        stage(job,'configured')
    log(job,'Save and reopen independent new project')
    u=browse(u,TEST);u=open_job(u,job);stage(job,'reopened')
    audit=audit_snapshot(job);log(job,'Native and inherited settings checks passed',checks=audit['check_count'])
    u,target=export(u,job);stage(job,'verified',backup=str(ROOT/'backups'/target.name),backup_sha256=sha(target),checks=audit['check_count'])
    log(job,'Project verified and full native backup complete',status='project_verified',completed=queue()['static_verified'])
    return u
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--start',type=int,default=1);ap.add_argument('--end',type=int,default=42);args=ap.parse_args();job=None
    try:
        policy()
        u=main_ui()
        for job in queue()['jobs']:
            if not args.start<=job['sequence']<=args.end or job['status']=='static_verified':continue
            u=build_zh(u,job) if job['language']=='ZH' else build_ja(u,job)
    except Exception as exc:
        if job:log(job,'Stopped for review',status='failed',error=str(exc),trace=traceback.format_exc())
        raise
