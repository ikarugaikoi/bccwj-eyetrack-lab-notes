# !__MASKED_LOCAL_PATH_0917__
'Assemble a self-contained Tobii input package; never creates native projects.\n\nXLSX must already have been authored by edit_tables.mjs. openpyxl is used here\nONLY to independently read and compare workbooks. All source paths are read-only.\n'
import argparse
from collections import Counter
from copy import deepcopy
import csv
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import zipfile

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
PROFILE = 'DELL238_FHD_56CM_FUSION120_20260908'
LAYOUT = {'MAIN': dict(zip('WHXY', [1, 1, 0, 0])),
          'DOT': dict(zip('WHXY', [.05, .07777778, .01, .06666667]))}
ORDERS = ['AB', 'BA', 'CD', 'DC']
GUIDE_CELLS = {'B2', 'B3', 'B4', 'B6', 'B13', 'B15', 'B16', 'B20'}


def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def read(p):
    return json.loads(p.read_text(encoding='utf-8'))


def write(p, value):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')


def copy(src, dst):
    if src.resolve() == dst.resolve():
        return
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dst)
    assert digest(src) == digest(dst)


def stem(i, order):
    return f'order{i:02d}_{order}_P_TRUE_Q_FALSE'


def compare_tables(baseline, out):
    reports=[]
    for i, order in enumerate(ORDERS, 1):
        a=baseline/'02_DESIGN_TABLES'/f'{stem(i, order)}.xlsx'
        b=out/'02_DESIGN_TABLES'/f'{stem(i, order)}_DELL1080.xlsx'
        old, new=load_workbook(a), load_workbook(b)
        assert old.sheetnames == new.sheetnames == ['Tobii_Design','Group_Map','Codebook']
        differences=[]
        features={}
        for name in old.sheetnames:
            x,y=old[name],new[name]
            assert (x.max_row,x.max_column)==(y.max_row,y.max_column)
            assert list(x.merged_cells.ranges)==list(y.merged_cells.ranges)
            assert x.freeze_panes==y.freeze_panes
            assert x.auto_filter.ref==y.auto_filter.ref
            tx=[(t.name,t.ref) for t in x.tables.values()]
            ty=[(t.name,t.ref) for t in y.tables.values()]
            assert tx==ty
            features[name]={'rows':y.max_row,'columns':y.max_column,'native_tables':ty,
                            'merged_cells_preserved':True,'freeze_and_filter_preserved':True}
            for row in x:
                for c in row:
                    v=y[c.coordinate]
                    assert v.data_type not in ['f','e'], (name,c.coordinate,'unexpected formula/error')
                    changed=c.value!=v.value or (c.value is not None and c.data_type!=v.data_type)
                    if not changed:
                        continue
                    main=name=='Tobii_Design' and c.row>1 and c.column in [5,7,9,11]
                    guide=name=='Codebook' and c.coordinate in GUIDE_CELLS
                    assert main or guide, (name,c.coordinate,c.value,v.value)
                    if main:
                        assert v.value == {5:.07777778,7:.06666667,9:1,11:0}[c.column]
                        assert c.data_type == v.data_type == 'n'
                    differences.append({'sheet':name,'cell':c.coordinate,'before':c.value,'after':v.value})
        count=73 if order in ['AB','BA'] else 64
        assert len(differences)==4*count+8
        reports.append({'order':order,'source':a.name,'output':b.name,
                        'source_sha256':digest(a),'output_sha256':digest(b),'data_rows':count,
                        'coordinate_cells_changed':count*4,'codebook_cells_changed':8,
                        'all_other_values_and_types_preserved':True,'sheet_features':features,
                        'differences':differences})
        old.close();new.close()
    return {'status':'PASS','authoring_engine':'@oai/artifact-tool',
            'independent_reader':'openpyxl read only; no workbook saves',
            'coordinate_cells_changed':1096,'codebook_cells_changed':32,
            'format_changes_intended':'DOT_H/DOT_Y show 8 decimals and wider columns; article_block and media_name widened to show identifiers; Codebook descriptions wrap, and changed or long descriptions have increased row height; remaining established table presentation retained.',
            'workbooks':reports}


def project_markdown(p):
    lines=[f"# {p['bilingual_project_name']}", '',
           '本文件是完整施工卡，不是已完成原生工程。中文/日文均须搭建。', '',
           f'- 配置：`{PROFILE}`；1920×1080，56cm目标，Fusion120 / 120Hz。',
           f"- 原逻辑名：`{p['base_project_name']}`；实际新建名：`{p['bilingual_project_name']}`。",
           '- 放独立DELL1080目录；如软件同名冲突可追加_DELL1080，并记录映射。',
           f"- 类型：{p['kind']}；条件：{p['order']}；半场：{p['half']}；文章：{' → '.join(p['article_blocks'])}。",
           f"- 导入：`{p['workbook_path']}`；表对象 `Tobii_Design`。",
           f"- 表文件 SHA256：`{p['workbook_sha256']}`。",
           f"- Group {p['group_count']}个；文章事件 {p['article_event_count']}次；含独立提示共{p['total_presented_image_instances']}次图像呈现（不含原生校准）。", '',
           '详读03_SPEC/REBUILD_REQUIREMENTS.md。MAIN=1/1/0/0；DOT=.05/.07777778/.01/.06666667；均Original。普通刺激100ms最短、Time=None；FIX无键；READ仅Space；QUESTION Q/P均可；结束仅7。', '',
           '## 顶层全部元素', '', '| 步骤 | 类型 | 对象名称 | 固定Source或内部Stimulus | 推进预设 |', '|---:|---|---|---|---|']
    for n,e in enumerate(p['elements'],1):
        if e['kind']=='calibration':
            name=e['name'];source='原生Calibration + Validation；9目标/Point/Timed';preset='CALIBRATION'
        elif e['kind']=='group':
            name=e['group_name'];source=e['stimulus_name'];preset=e['preset']
        else:
            name=e['stimulus_name'];source=Path(e['media_path']).stem;preset=e['preset']
        lines.append(f"| {n} | {e['kind']} | `{name}` | `{source}` | {preset} |")
    lines+=['','## 各文章组绑定与实际展开行','','每组只建一个内部模板，不为每张阅读图另造组。两个Subset的Set at recording start关闭；无额外operator。', '']
    for e in p['elements']:
        if e['kind']!='group':continue
        lines += [f"### {e['group_name']}", '',f"- 内部Stimulus：`{e['stimulus_name']}`。",
                  f"- Subset 1：`article_block={e['subsets'][0]['values'][0]}`；Subset 2：`row_type={e['subsets'][1]['values'][0]}`。",
                  f"- 预设：`{e['preset']}`。"]
        for c in e['containers']:
            bindings=', '.join(k+'→'+v for k,v in c['layout_bindings'].items())
            lines.append(f"- 容器`{c['name']}`：Source {c['source_mode']}=`{c['source_value']}`；坐标绑定{bindings}。")
        lines+=['','| Excel行 | event_id | sample_screen | media_name |','|---:|---|---|---|']
        lines += [f"| {ev['excel_row']} | `{ev['event_id']}` | `{ev['sample_screen']}` | `{ev['media_name']}` |" for ev in e['expected_events']]
        lines.append('')
    lines+=['## 本工程验收','','- [ ] 新版本/语言/表哈希/所有名称和Subset正确。',
            '- [ ] 固定说明页MAIN手填1/1/0/0；文章容器绑定正确；所有Source解析。',
            '- [ ] Time=None；Q/P两键回答；7结束；仅DOT有Continuous300ms/数据丢失重置34ms。',
            '- [ ] 校准恰好一次；所有正文页/题/结束按以上顺序展开。',
            '- [ ] 保存、重新打开、只读原生检查、现场试跑和导出分别记录，不把静态通过当现场通过。','']
    return '\n'.join(lines)


def assemble(args):
    baseline=args.baseline.resolve();audit=args.audit.resolve();out=args.out.resolve()
    assets=(args.assets or baseline).resolve()
    assert out!=baseline and out!=audit
    #  Prove the untouched full legacy package when its full manifest is available.
    original_manifest=read(baseline/'HASHES_SHA256.json') if (baseline/'HASHES_SHA256.json').exists() else None
    if original_manifest:
        for rel,meta in original_manifest['files'].items():
            assert digest(baseline/rel)==meta['sha256'],rel
    out.mkdir(parents=True,exist_ok=True)
    for p in (out/'02_DESIGN_TABLES').glob('*.inspect.ndjson'):
        target=out/'06_REFERENCE_DO_NOT_IMPORT/TABLE_QA'/p.name
        target.parent.mkdir(parents=True,exist_ok=True)
        shutil.move(p,target)
    diff=compare_tables(baseline,out)
    write(out/'06_REFERENCE_DO_NOT_IMPORT/TABLE_DIFF_REPORT.json',diff)
    media_manifest=[]
    for p in sorted((assets/'01_MEDIA').rglob('*.png')):
        rel=p.relative_to(assets);copy(p,out/rel)
        media_manifest.append({'path':rel.as_posix(),'sha256':digest(p),'bytes':p.stat().st_size})
    assert len(media_manifest)==131
    for p in (assets/'06_REFERENCE_DO_NOT_IMPORT/AOI').rglob('*.csv'):
        copy(p,out/p.relative_to(assets))
    for rel in ['BUILD_SELECTION.json','03_SPEC/INSTRUCTION_TEXTS.json','03_SPEC/QUESTION_SCORING_P_TRUE_Q_FALSE.json',
                '03_SPEC/RECOVERY_POLICY.md','03_SPEC/RECOVERY_TEXT_REVIEW.md','05_QA/LANGUAGE_COPY_CHECKLIST.md',
                '05_QA/BUILD_PROGRESS_TEMPLATE.json','05_QA/INCIDENT_LOG_TEMPLATE.json']:
        copy(baseline/rel,out/rel)
    for name in ['00_START_HERE.md','00_WINDOWS_CODEX_PROMPT.md']:
        copy(HERE/name,out/name)
    copy(HERE/'REBUILD_REQUIREMENTS.md',out/'03_SPEC/REBUILD_REQUIREMENTS.md')
    params=read(baseline/'03_SPEC/BUILD_PARAMETERS.json')
    params.update({'profile_id':PROFILE,'layout':LAYOUT,'presentation_resolution':[1920,1080],
                   'assumed_display_mm':[527,296],'actual_display_verified':False,
                   'eye_midpoint_to_screen_plane_target_mm':560,
                   'tracker_model':'Tobii Pro Fusion 120','tracker_model_confirmed_by_lab':True,
                   'tracker_confirmation_source':'User relayed laboratory warranty hardware designation, 2026-09-08',
                   'recording_frequency_target_hz':120,'actual_recording_frequency_verified':False,
                   'continuous_gaze_semantics':'An effective valid gaze outside restarts continuous dwell; 34 ms reset is specifically data loss, not valid off-target grace.'})
    write(out/'03_SPEC/BUILD_PARAMETERS.json',params)
    geometry_old=read(audit/'PROPOSED_GEOMETRY.json')
    geometry={k:deepcopy(v) for k,v in geometry_old.items() if k.startswith(('distance_','angular_'))}
    geometry.update({'profile_id':PROFILE,'assumed_display_mm':[527,296],'resolution_px':[1920,1080],
                     'physical_display_verified':False,'normalized':LAYOUT,'coordinate_rounding_decimals':8,
                     'exact_dot_fractions':{'H':'84/1080','Y':'72/1080'},
                     'main_xywh_px':[0,0,1920,1080],'dot_xywh_px':[19.2,72,96,84],
                     'dot_visible_alpha_bbox_local':[40,34,56,50],'dot_visible_center_screen_px':[67.2,114],
                     'screen_to_image_offset_px':[0,0],
                     'note':'Pixel rectangles are ideal targets; 8 decimal normalized values differ by less than 0.00001 px. Onsite verify native render rectangles. Historical 1200p data have a different transform.'})
    write(out/'03_SPEC/DISPLAY_GEOMETRY.json',geometry)
    all_projects=[];catalog=['# 84个新工程索引','','ZH先42个，再对应复制JA42个。此清单不是施工完成记录。配置 '+PROFILE+'。','','| 序号 | 逻辑工程 | 类型 | 文章 | 中文卡 | 日文卡 |','|---:|---|---|---|---|---|']
    for lang in ['ZH','JA']:
        spec=read(baseline/f'04_PROJECT_BLUEPRINTS/PROJECTS_{lang}.json')
        spec.update({'schema_version':2,'profile_id':PROFILE,'presets':params['presets']})
        for index,p in enumerate(spec['projects'],1):
            old_path=p['workbook_path']
            p['workbook_path']=old_path.replace('.xlsx','_DELL1080.xlsx')
            p['workbook_sha256']=digest(out/p['workbook_path'])
            p['presentation_resolution']=[1920,1080]
            p['profile_id']=PROFILE
            p['native_build_status']='not_built_in_this_package'
            for e in p['elements']:
                for c in e.get('containers',[]):
                    c['layout']=deepcopy(LAYOUT[c['name']])
            md=out/f"04_PROJECT_BLUEPRINTS/{lang}/{index:02d}_{p['base_project_name']}.md"
            md.parent.mkdir(parents=True,exist_ok=True);md.write_text(project_markdown(p),encoding='utf-8')
            if lang=='ZH':
                f=f"{index:02d}_{p['base_project_name']}.md"
                catalog.append(f"| {index} | `{p['base_project_name']}` | {p['kind']} | {' → '.join(p['article_blocks'])} | [ZH](../04_PROJECT_BLUEPRINTS/ZH/{f}) | [JA](../04_PROJECT_BLUEPRINTS/JA/{f}) |")
            all_projects.append(p)
        write(out/f'04_PROJECT_BLUEPRINTS/PROJECTS_{lang}.json',spec)
    (out/'03_SPEC/PROJECT_CATALOG.md').write_text('\n'.join(catalog)+'\n',encoding='utf-8')
    plan=read(baseline/'04_PROJECT_BLUEPRINTS/SELECTED_BUILD_PLAN.json')
    plan.update({'date':'2026-09-08','profile_id':PROFILE,'native_projects_already_built':0})
    for j in plan['jobs']:
        j['workbook_path']=j['workbook_path'].replace('.xlsx','_DELL1080.xlsx')
        j['workbook_sha256']=digest(out/j['workbook_path'])
        j['profile_id']=PROFILE
    write(out/'04_PROJECT_BLUEPRINTS/SELECTED_BUILD_PLAN.json',plan)
    progress=read(out/'05_QA/BUILD_PROGRESS_TEMPLATE.json');progress['profile_id']=PROFILE
    write(out/'05_QA/BUILD_PROGRESS_TEMPLATE.json',progress)
    scope=[]
    for p in all_projects:
        scope.append({'project':p['bilingual_project_name'],'language':p['language'],'kind':p['kind'],
                      'order':p['order'],'half':p['half'],'articles':'|'.join(p['article_blocks']),
                      'groups':p['group_count'],'status':'NOT_BUILT_NOT_VERIFIED'})
    with (out/'05_QA/PROJECT_SCOPE.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(scope[0]));w.writeheader();w.writerows(scope)
    expected={'notice':'QA-only, not a Pro Lab import. All element layouts and top-level geometry describe the same NEW profile. No native/hardware verification has occurred.',
              'profile_id':PROFILE,'geometry':LAYOUT,
              'projects':[{'project':p['bilingual_project_name'],
                           'source_blueprint':f"04_PROJECT_BLUEPRINTS/PROJECTS_{p['language']}.json",
                           'elements':p['elements'],'presets':params['presets'],'expected_display':[1920,1080]} for p in all_projects]}
    write(out/'05_QA/EXPECTED_PROJECTS_QA_ONLY.json',expected)
    verifier=(baseline/'05_QA/verify_package.py').read_text(encoding='utf-8')
    swaps={'[1, .9, 0, .05]':'[1, 1, 0, 0]', '[.05, .07, .01, .11]':'[.05, .07777778, .01, .06666667]',
           'P_TRUE_Q_FALSE.xlsx':'P_TRUE_Q_FALSE_DELL1080.xlsx', '[1920, 1200]':'[1920, 1080]',
           "builder's openpyxl reader":"artifact-tool author's output (independent XML reader)"}
    for before,after in swaps.items():
        assert before in verifier,before
        verifier=verifier.replace(before,after)
    (out/'05_QA/verify_package.py').write_text(verifier,encoding='utf-8')
    for name in ['check_native_project.py','read_tracker_config.py']:
        copy(audit/name,out/'05_QA'/name)
    onsite=(audit/'03_ONSITE_CHECKLIST.md').read_text(encoding='utf-8')
    onsite=onsite.replace('## B. 工程迁移与图像','## B. 新建工程与图像').replace('四列新表已更新到原表对象，绑定未丢','已导入本完整包的_DELL1080表，所有绑定已核对')
    (out/'05_QA/ONSITE_CHECKLIST.md').write_text(onsite,encoding='utf-8')
    ret=(audit/'04_RETURN_PACKAGE.md').read_text(encoding='utf-8').replace('py -3 read_tracker_config.py','py -3 05_QA\\read_tracker_config.py')
    (out/'05_QA/RETURN_PACKAGE.md').write_text(ret,encoding='utf-8')
    site=read(audit/'SITE_LOG_TEMPLATE.json');site['proposed_profile']['profile_id']=PROFILE
    site['proposed_profile']['tracker_hardware_model_confirmed']='Tobii Pro Fusion 120'
    write(out/'05_QA/SITE_LOG_TEMPLATE.json',site)
    for name in ['ACCEPTANCE_CHECKLIST.md','CHECKER_USAGE.md']:
        copy(HERE/name,out/'05_QA'/name)
    if args.qa:
        for p in args.qa.iterdir():
            if p.is_file() and p.suffix in ['.png','.json']:
                copy(p,out/'06_REFERENCE_DO_NOT_IMPORT/TABLE_QA'/p.name)
    write(out/'06_REFERENCE_DO_NOT_IMPORT/MEDIA_PROVENANCE.json',
          {'source':'Previously user-approved 2026-09-05 package','unchanged_images':True,'count':131,'files':media_manifest})
    write(out/'06_REFERENCE_DO_NOT_IMPORT/SOURCE_INPUT_MANIFEST.json',
          {'baseline_local_path':str(baseline),'assets_local_path':str(assets),'audit_local_path':str(audit),
           'paths_are_provenance_not_windows_dependencies':True,'source_full_manifest':original_manifest})
    #  Include sufficient metadata and workbook baselines to reproduce this authoring
    #  process locally; reused media/AOI remain in normal package dirs, not duplicated.
    snap=out/'08_BUILD_SOURCE/baseline_inputs'
    baseline_paths=['BUILD_SELECTION.json','03_SPEC/BUILD_PARAMETERS.json','03_SPEC/INSTRUCTION_TEXTS.json',
                    '03_SPEC/QUESTION_SCORING_P_TRUE_Q_FALSE.json','03_SPEC/RECOVERY_POLICY.md','03_SPEC/RECOVERY_TEXT_REVIEW.md',
                    '04_PROJECT_BLUEPRINTS/PROJECTS_ZH.json','04_PROJECT_BLUEPRINTS/PROJECTS_JA.json',
                    '04_PROJECT_BLUEPRINTS/SELECTED_BUILD_PLAN.json','05_QA/verify_package.py',
                    '05_QA/LANGUAGE_COPY_CHECKLIST.md','05_QA/BUILD_PROGRESS_TEMPLATE.json','05_QA/INCIDENT_LOG_TEMPLATE.json']
    baseline_paths += [f'02_DESIGN_TABLES/{stem(i, o)}.xlsx' for i,o in enumerate(ORDERS,1)]
    for rel in baseline_paths:copy(baseline/rel,snap/rel)
    for name in ['PROPOSED_GEOMETRY.json','angular_geometry.py','03_ONSITE_CHECKLIST.md','04_RETURN_PACKAGE.md',
                 'SITE_LOG_TEMPLATE.json','check_native_project.py','read_tracker_config.py']:
        copy(audit/name,out/'08_BUILD_SOURCE/audit_inputs'/name)
    for p in HERE.iterdir():
        if p.is_file() and p.suffix in ['.py','.mjs','.md']:
            copy(p,out/'08_BUILD_SOURCE'/p.name)
    # Historical image-generator source is excluded from this distribution.
    # Approved media are reused above; do not collect old source-tree copies.
    (out/'07_BUILD_LOGS').mkdir(exist_ok=True)
    (out/'07_BUILD_LOGS/README.md').write_text('# Windows实际施工日志\n\n复制05_QA的空白模板后填写，不修改模板本身。记录实际路径、版本、完成/失败/待验证、证据。此目录初始没有已完成原生工程。\n',encoding='utf-8')
    counter=Counter(e['kind'] for p in all_projects for e in p['elements'])
    containers=[c for p in all_projects for e in p['elements'] for c in e.get('containers',[])]
    assert len(all_projects)==84 and counter['group']==756 and len(containers)==1292
    assert sum(len(j['instruction_replacements']) for j in plan['jobs'])==142
    if original_manifest:
        for rel,meta in original_manifest['files'].items():
            assert digest(baseline/rel)==meta['sha256'],'Source changed: '+rel
    report={'status':'MAC_INPUTS_VALIDATED_NATIVE_BUILD_PENDING','profile_id':PROFILE,
            'workbooks':4,'coordinate_cells_changed':1096,'codebook_cells_changed':32,
            'media_png':131,'original_media_byte_identical':True,'projects':84,'groups':756,
            'ordinary_templates':counter['group']+counter['image_stimulus'],'containers':len(containers),
            'instruction_replacements_ja':142,'old_package_hashes_unchanged':bool(original_manifest),
            'native_projects_built':0,'native_projects_verified':False,'hardware_verified':False}
    write(out/'06_REFERENCE_DO_NOT_IMPORT/PACKAGE_BUILD_CHECKS.json',report)
    refresh_manifest(out)
    result=subprocess.run([sys.executable,str(out/'05_QA/verify_package.py')],check=True,capture_output=True,text=True)
    write(out/'06_REFERENCE_DO_NOT_IMPORT/PACKAGE_VERIFY_RESULT.json',json.loads(result.stdout))
    refresh_manifest(out)
    print(json.dumps(report,ensure_ascii=False,indent=2))


def refresh_manifest(out):
    files={p.relative_to(out).as_posix():{'bytes':p.stat().st_size,'sha256':digest(p)}
           for p in sorted(out.rglob('*')) if p.is_file() and p.name!='HASHES_SHA256.json'
           and '07_BUILD_LOGS' not in p.relative_to(out).parts and '__pycache__' not in p.parts}
    write(out/'HASHES_SHA256.json',{'date':'2026-09-08','profile_id':PROFILE,'algorithm':'sha256',
                                 'excluded':['HASHES_SHA256.json','07_BUILD_LOGS/*','**/__pycache__/*'],'files':files})


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--baseline',required=True,type=Path)
    parser.add_argument('--audit',required=True,type=Path)
    parser.add_argument('--out',required=True,type=Path)
    parser.add_argument('--assets',type=Path)
    parser.add_argument('--qa',type=Path)
    assemble(parser.parse_args())
