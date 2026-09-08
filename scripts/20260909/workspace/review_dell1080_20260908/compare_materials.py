import hashlib, json, re
from pathlib import Path
from collections import Counter
import openpyxl

ROOT = Path(__file__).resolve().parent
NEW = ROOT / 'BCCWJ_TOBII_FULL_20260908_DELL1080_56CM_120HZ'
OLD = Path('__MASKED_LOCAL_PATH_0290__')
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p): return json.loads(p.read_text(encoding='utf-8-sig'))
def diff(a,b,path=''):
    if isinstance(a,dict) and isinstance(b,dict):
        out=[]
        for k in sorted(a.keys()|b.keys()):
            p=path+'/'+str(k)
            if k not in a or k not in b: out.append({'path':p,'old':a.get(k),'new':b.get(k)})
            else: out.extend(diff(a[k],b[k],p))
        return out
    if isinstance(a,list) and isinstance(b,list) and len(a)==len(b):
        return [d for i,(x,y) in enumerate(zip(a,b)) for d in diff(x,y,path+'/'+str(i))]
    return [] if a==b else [{'path':path,'old':a,'new':b}]

report={'zip_sha256':sha(Path('__MASKED_LOCAL_PATH_0295__'))}
media=list((NEW/'01_MEDIA').rglob('*.png'))
report['media']={'count':len(media),'byte_identical':0,'different':[],'missing_old':[]}
for p in media:
    old=OLD/p.relative_to(NEW)
    if not old.exists(): report['media']['missing_old'].append(str(p.relative_to(NEW)))
    elif sha(p)==sha(old): report['media']['byte_identical']+=1
    else: report['media']['different'].append(str(p.relative_to(NEW)))
aois=list((NEW/'06_REFERENCE_DO_NOT_IMPORT/AOI').rglob('*.csv'))
report['aoi']={'count':len(aois),'byte_identical':sum((OLD/p.relative_to(NEW)).exists() and sha(p)==sha(OLD/p.relative_to(NEW)) for p in aois)}
report['tables']=[]
for p in sorted((NEW/'02_DESIGN_TABLES').glob('*.xlsx')):
    old=OLD/'02_DESIGN_TABLES'/p.name.replace('_DELL1080','')
    a=openpyxl.load_workbook(old,read_only=True,data_only=False)
    b=openpyxl.load_workbook(p,read_only=True,data_only=False)
    entry={'file':p.name,'same_sheet_names':a.sheetnames==b.sheetnames,'sheets':[]}
    for name in b.sheetnames:
        ar=list(a[name].values);br=list(b[name].values)
        changes=[]
        for i in range(max(len(ar),len(br))):
            av=ar[i] if i<len(ar) else ();bv=br[i] if i<len(br) else ()
            for j in range(max(len(av),len(bv))):
                x=av[j] if j<len(av) else None;y=bv[j] if j<len(bv) else None
                if x!=y:changes.append({'cell':f'{openpyxl.utils.get_column_letter(j + 1)}{i + 1}','column':str(br[0][j]) if j<len(br[0]) else str(j),'old':x,'new':y})
        entry['sheets'].append({'sheet':name,'old_rows':len(ar),'new_rows':len(br),'old_columns':a[name].max_column,'new_columns':b[name].max_column,'change_count':len(changes),'by_column':dict(Counter(x['column'] for x in changes)),'changes':changes})
    a.close();b.close();report['tables'].append(entry)
report['json']={}
for rel in ['03_SPEC/BUILD_PARAMETERS.json','03_SPEC/INSTRUCTION_TEXTS.json','03_SPEC/QUESTION_SCORING_P_TRUE_Q_FALSE.json','04_PROJECT_BLUEPRINTS/PROJECTS_ZH.json','04_PROJECT_BLUEPRINTS/PROJECTS_JA.json','04_PROJECT_BLUEPRINTS/SELECTED_BUILD_PLAN.json']:
    changes=diff(read(OLD/rel),read(NEW/rel))
    report['json'][rel]={'count':len(changes),'by_path':dict(Counter(re.sub('/\\d+(?=/|$)','/*',x['path']) for x in changes)),'changes':changes}
(ROOT/'material_comparison.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({**{k:v for k,v in report.items() if k not in ('tables','json')},'tables':[{**e,'sheets':[{k:v for k,v in s.items() if k!='changes'} for s in e['sheets']]} for e in report['tables']],'json':{k:{f:v for f,v in e.items() if f!='changes'} for k,e in report['json'].items()}},ensure_ascii=False,indent=2))
